#!/usr/bin/env python3
"""Percepciones del DANE con geografía fina → celdas agregadas para anclar el sondeo.

  ECV 2025 (anexos, cuadros 35-39): opinión del jefe/a del hogar por DEPARTAMENTO ×
      cabecera/resto: ingresos, si se considera pobre, seguridad del barrio, situación
      económica hace y dentro de 12 meses.
  Pulso Social 2023 (anexos abr-may-jun, promedio): opinión por CIUDAD (23 capitales y
      sus áreas, Sincelejo incluida), sexo y edad.

Solo agregados publicados por el DANE (anexos públicos, no microdatos). Salida:
  simcolombia/data/percepcion_items.json   todos los ítems crudos (para etiquetar)
  web/opinion/ECV_*.json, web/opinion/EPS_*.json   celdas en el formato de votoCelda

  uv run --with openpyxl python simcolombia/pipeline/fuentes_dane_percepcion.py
"""
import json
import re
import unicodedata
from collections import defaultdict
from pathlib import Path

import openpyxl

RAIZ = Path(__file__).resolve().parents[2]
RAW = RAIZ / "simcolombia" / "data" / "raw_v2"
OUT = RAIZ / "web" / "opinion"
MARG = json.loads((RAIZ / "web" / "marginals.json").read_text())["departamentos"]


def norm(s):
    s = unicodedata.normalize("NFD", str(s or "")).encode("ascii", "ignore").decode().lower()
    return re.sub(r"[^a-z ]", "", s).strip()


DPTO_DE = {norm(v["nombre"]).replace(" dc", ""): k for k, v in MARG.items()}
DPTO_DE.update({"bogota": "11", "bogota dc": "11", "san andres": "88", "archipielago de san andres": "88",
                "valle": "76", "norte de santander": "54", "la guajira": "44", "guajira": "44"})


def dpto(nombre):
    n = norm(nombre)
    if n in DPTO_DE:
        return DPTO_DE[n]
    for k, v in DPTO_DE.items():
        if n.startswith(k) or k.startswith(n):
            return v
    return None


def filas(ws):
    return [list(r) for r in ws.iter_rows(values_only=True)]


# ── ECV 2025 ────────────────────────────────────────────────────────────────
ECV_CUADROS = {35: "ingresos", 36: "pobre", 37: "barrio", 38: "economia_antes", 39: "economia_despues"}


def ecv():
    wb = openpyxl.load_workbook(RAW / "anex-ECV-2025.xlsx", read_only=True, data_only=True)
    items = {}
    for n, clave in ECV_CUADROS.items():
        F = filas(wb[f"Cuadro {n}"])
        titulo = next((str(c) for r in F[:12] for c in r if isinstance(c, str) and len(c) > 30), f"Cuadro {n}")
        h = next(i for i, r in enumerate(F) if r and str(r[0] or "").strip().startswith("Departamento"))
        cats = [str(c).strip() for c in F[h][3:] if c]          # después de Departamento, Área, Total hogares
        if len(cats) < 2:                                       # 38 y 39: las categorías van en la fila de abajo
            h += 1
            cats = [str(c).strip() for c in F[h] if c]
        # cada fila, sin las celdas vacías (el DANE mete columnas separadoras en la
        # mitad de la hoja): total(4) y por categoría 8 = Total, LI, LS, CVE, %, LI, LS, CVE
        cel, dep = {}, None
        for r in F[h + 2:]:
            if not r or r[1] is None:
                continue
            if r[0]:
                dep = "nacional" if norm(r[0]).startswith("total nacional") else dpto(r[0])
            area = {"total": "total", "cabecera": "cabecera"}.get(norm(r[1]), "resto")
            if dep is None:
                continue
            v = [x for x in r[2:] if x is not None]
            try:
                d = {str(i + 1): round(float(v[8 + 8 * i]) / 100, 4) for i in range(len(cats))}
            except (TypeError, ValueError, IndexError):
                continue
            if abs(sum(d.values()) - 1) > .03:                  # si no suma 100 %, la fila está mal leída: fuera
                continue
            cel[f"{dep}|{area}"] = d
        items[f"ECV:{clave}"] = {"fuente": "DANE ECV 2025", "cuadro": n, "titulo": titulo,
                                 "opciones": {str(i + 1): c for i, c in enumerate(cats)}, "celdas": cel}
    return items


# ── Pulso Social 2023 ───────────────────────────────────────────────────────
def pulso():
    acum = defaultdict(lambda: defaultdict(lambda: defaultdict(list)))
    meta = {}
    for mes in ("abr", "may", "jun"):
        wb = openpyxl.load_workbook(RAW / f"anex-EPS-{mes}2023.xlsx", read_only=True, data_only=True)
        for hoja in wb.sheetnames:
            if hoja in ("Índice", "ICC"):
                continue
            F = filas(wb[hoja])
            preg = next((re.split(r"\n\s*Totales y", str(c))[0].replace("\n", " ").strip()
                         for r in F[:6] for c in r if isinstance(c, str) and re.match(r"^\w+\.", c)), None)
            if not preg:
                continue
            bloque, ncat, etiqueta = None, None, ""
            tipo = lambda c0: ("sexo" if c0.startswith("Sexo") else "edad" if c0.startswith("Edad")
                               else "ciudad" if (c0.startswith("Ciudad") or "ciudades" in c0) and "Total" not in c0 else None)
            for i, r in enumerate(F):
                if not r:
                    continue
                c0 = str(r[0]).strip() if r[0] is not None else ""
                if r[1] and str(r[1]).strip() in ("Total personas", "%"):
                    continue                           # subencabezado "Total personas | %": no es de categorías
                if len(r) > 2 and r[1] and isinstance(r[1], str) and not isinstance(r[2], (int, float)):
                    # encabezado de bloque: "Sexo | Muy seguro/a | cve | Ls | Li | …"; a veces la
                    # etiqueta (Sexo, Ciudad) va sola en la fila de arriba y las categorías abajo
                    cats = [str(x).strip() for x in r[1:] if x and str(x).strip() not in ("cve", "Ls", "Li", "Total")]
                    bloque = tipo(c0 or etiqueta)
                    ncat = len(cats)
                    continue
                if c0 and all(x is None for x in r[1:]):
                    etiqueta = c0                      # "Sexo" solo: el encabezado viene en la fila siguiente
                    continue
                if not c0:
                    continue
                if bloque is None and not c0.startswith("Total 23"):
                    continue
                nums = [x for x in r[1:] if isinstance(x, (int, float))]
                if not ncat or len(nums) < ncat * 5:
                    continue
                props = [nums[5 * k + 1] for k in range(ncat)]      # [total, %, cve, Ls, Li] por categoría
                if not all(0 <= p <= 1.0001 for p in props) or sum(props) < .9:
                    continue
                if c0.startswith("Total 23"):
                    dim, val = "*", "*"
                elif bloque == "sexo":
                    dim, val = "sexo", {"hombres": "1", "mujeres": "2"}.get(norm(c0))
                elif bloque == "edad":
                    dim, val = "ge3", {"10 a 24 anos": "0", "25 a 54 anos": "1", "55 anos o mas": "2"}.get(norm(c0))
                elif bloque == "ciudad":
                    dim, val = "ciudad", re.sub(r"\s*am$", "", norm(c0)).strip()   # "BucaramangaAM" → bucaramanga (solo el sufijo)
                else:
                    continue
                if val is None:
                    continue
                acum[hoja][dim][val].append(props)
                meta[hoja] = {"pregunta": preg, "opciones": cats}
    items = {}
    for hoja, dims in acum.items():
        m = meta[hoja]
        prom = lambda L: {str(k + 1): round(sum(x[k] for x in L) / len(L), 4) for k in range(len(L[0]))}
        D = {"*": prom(dims["*"]["*"])} if "*" in dims else {}
        for dim in ("sexo", "ge3", "ciudad"):
            if dim in dims:
                D[dim] = {v: prom(L) for v, L in dims[dim].items()}
        if "*" in D and "ciudad" in D:
            items[f"EPS:{hoja}"] = {"fuente": "DANE Pulso Social 2023 (abr-jun)", "pregunta": m["pregunta"],
                                    "opciones": {str(i + 1): c for i, c in enumerate(m["opciones"])}, "celdas": D}
    return items


def main():
    E, P = ecv(), pulso()
    todo = {**E, **P}
    (RAIZ / "simcolombia" / "data" / "percepcion_items.json").write_text(json.dumps(todo, ensure_ascii=False, indent=1))
    print(f"ECV: {len(E)} ítems · Pulso: {len(P)} ítems con ciudades")
    s = E["ECV:pobre"]["celdas"]
    print("ECV 'se considera pobre' Sucre:", s.get("70|total"), "cabecera", s.get("70|cabecera"), "resto", s.get("70|resto"), "· nacional", s.get("nacional|total"))
    b = P.get("EPS:bs10")
    if b:
        print("Pulso bs10 (seguro de noche) ciudades:", sorted(b["celdas"]["ciudad"])[:30])
        print("  Sincelejo:", b["celdas"]["ciudad"].get("sincelejo"), "· 23 ciudades:", b["celdas"]["*"])


if __name__ == "__main__" and "--publicar" not in __import__("sys").argv:
    main()


# ── publicar: celdas en web/opinion y entradas en el banco del sondeo ──────────
def publicar():
    """Lee data/percepcion_items.json + data/percepcion_etiquetas.json (etiquetado por el
    enjambre GLM y corregido a mano) y escribe:
      web/opinion/ECV_*.json  {"*": nacional, "dc": {"70|cabecera": …, "70|resto": …}}
      web/opinion/EPS_*.json  {"*": 23 ciudades, "sexo", "ge3", "ciudad": {"sincelejo": …}}
      web/banco.json          reemplaza las entradas ECV:/EPS: anteriores."""
    items = json.loads((RAIZ / "simcolombia" / "data" / "percepcion_items.json").read_text())
    lab = json.loads((RAIZ / "simcolombia" / "data" / "percepcion_etiquetas.json").read_text())
    banco = [b for b in json.loads((RAIZ / "web" / "banco.json").read_text()) if not b["codigo"].startswith(("ECV:", "EPS:"))]
    n = 0
    for cod, it in items.items():
        L = lab.get(cod)
        if not L or not L.get("usar"):
            continue
        c = it["celdas"]
        if cod.startswith("ECV:"):
            D = {"*": c["nacional|total"], "dc": {k: v for k, v in c.items() if not k.startswith("nacional") and not k.endswith("|total")}}
            fuente = "DANE ECV 2025 (por departamento y zona)"
        else:
            D = {"*": c["*"], **{d: c[d] for d in ("sexo", "ge3") if d in c},
                 "ciudad": {k: v for k, v in c["ciudad"].items() if k != "total"}}
            fuente = "DANE Pulso Social 2023 (23 ciudades)"
        (OUT / f"{cod.replace(':', '_')}.json").write_text(json.dumps(D, ensure_ascii=False, separators=(",", ":")))
        banco.append({"codigo": cod, "fuente": fuente, "texto": L["texto"], "opciones": it["opciones"],
                      "afirmativas": L.get("afirmativas", []), "sustantivas": L.get("sustantivas", list(it["opciones"]))})
        n += 1
    (RAIZ / "web" / "banco.json").write_text(json.dumps(banco, ensure_ascii=False, separators=(",", ":")))
    print(f"publicados {n} ítems · banco: {len(banco)} preguntas")


if __name__ == "__main__" and "--publicar" in __import__("sys").argv:
    publicar()
